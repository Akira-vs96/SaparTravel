"""Correct the RAID-linked financial workbook while retaining the historical LZ8 package."""
import argparse
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

ROOT = Path(__file__).resolve().parents[2]


def raid_total(text):
    values = []
    for line in text.splitlines():
        if re.match(r"\| R-\d\d \|", line):
            cells = [c.strip() for c in line.strip("|").split("|")]
            p, impact, expected = [Decimal(cells[i].replace(",", ".")) for i in (2, 3, 4)]
            if p * impact != expected:
                raise ValueError("RAID row EMV does not match probability times impact")
            values.append(expected)
    if len(values) < 20:
        raise ValueError("At least twenty risks are required")
    total = sum(values)
    declared = re.search(r"\*\*Итого EMV:\*\*\s*([\d,]+)", text)
    if not declared or Decimal(declared[1].replace(",", ".")) != total:
        raise ValueError("Declared RAID total differs from the actual sum")
    return total


def build(directory):
    from openpyxl import load_workbook
    source = ROOT / "docs/lab-8/SaparTravel_LZ8_Financial_Model.xlsx"
    emv = raid_total((ROOT / "docs/security/raid.md").read_text())
    directory.mkdir(parents=True, exist_ok=True)
    book = load_workbook(source, data_only=False)
    inputs = book["Inputs"]
    inputs["B17"] = float(emv)
    reserve_row = inputs.max_row + 1
    inputs.append(["Резерв, уже исключённый из net capacity WBS", 36, "ч", "docs/plan/wbs.md; не вычитать дважды"])
    gap = max(Decimal(0), Decimal(171) + emv + 16 - (204 + 36))
    book["Checks"]["B10"] = f"=MAX(0,Inputs!B19+Inputs!B17+Inputs!B18-(Inputs!B20+Inputs!B{reserve_row}))"
    book["Checks"]["C10"] = float(gap)
    book["Guide"]["B2"] = "1.2 / 30.10.2026; исправлены сумма EMV и двойной учёт capacity"
    book["Guide"]["B8"] = f"Пересчитать формулы; Checks совпадает с Expected, дефицит capacity = {gap} ч"
    book["Guide"]["B10"] = f"EMV {emv} ч; management 16 ч; 36 ч уже исключены из net capacity"
    target = directory / source.name
    book.save(target)
    with tempfile.TemporaryDirectory(prefix="sapar-finance-") as temp:
        temp = Path(temp)
        result = subprocess.run(["libreoffice", "-env:UserInstallation=" + (temp / "profile").as_uri(),
                                 "--headless", "--convert-to", "xlsx", "--outdir", str(temp / "recalculated"), str(target)],
                                capture_output=True, timeout=120)
        recalculated = temp / "recalculated" / target.name
        if result.returncode or not recalculated.is_file():
            raise RuntimeError("LibreOffice recalculation failed; no financial result certified")
        shutil.copyfile(recalculated, target)
    values = load_workbook(target, data_only=True)
    checks = []
    for row in values["Checks"].iter_rows(min_row=2):
        name, actual, expected = [c.value for c in row[:3]]
        if name is not None:
            checks.append({"name": name, "actual": actual, "expected": expected,
                           "passed": isinstance(actual, (int, float)) and abs(actual - expected) < .00001})
    errors = [(sheet.title, c.coordinate) for sheet in values for row in sheet for c in row if c.data_type == "e"]
    report = {"version": "1.2", "as_of": "2026-10-30", "status": "forecast, not actual expenses",
              "currency": "KZT", "fx_reference_date": "2026-10-23", "risk_emv_hours": float(emv),
              "capacity_gap_hours": float(gap), "original_risk_hours": 63.2,
              "original_capacity_gap_hours": 46.2,
              "annual_base_kzt": values["Budget12"]["N12"].value,
              "risk_reserve_kzt": values["Budget12"]["N17"].value,
              "baseline_with_risk_kzt": values["Budget12"]["N18"].value,
              "management_reserve_kzt": values["Budget12"]["N19"].value,
              "funding_limit_kzt": values["Budget12"]["N20"].value,
              "actual_costs": None, "checks": checks, "spreadsheet_errors": errors,
              "source_workbook_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
              "passed": all(c["passed"] for c in checks) and not errors}
    (directory / "finance-validation.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    if not report["passed"]:
        raise RuntimeError("Corrected workbook failed independent arithmetic checks")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-directory", type=Path, default=ROOT / "docs/labs/lab-08/evidence")
    args = parser.parse_args()
    result = build(args.output_directory.resolve())
    print(json.dumps({k: result[k] for k in ("risk_emv_hours", "capacity_gap_hours", "funding_limit_kzt", "passed")}))


if __name__ == "__main__":
    main()
