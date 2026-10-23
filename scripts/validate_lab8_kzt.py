"""Independently audit the LZ8 package expressed in Kazakhstan tenge.

Run after recalculating the workbook and regenerating PDF exports:
    .venv/bin/python scripts/validate_lab8_kzt.py

Provider evidence retains its original USD denomination.  Project budgets,
unit economics, forecasts and comparison outputs must be expressed in KZT.
"""
from __future__ import annotations

import csv
import json
import math
import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "lab-8"
MAU = [0, 0, 0, 100, 200, 400, 600, 800, 1000, 1300, 1600, 2000]
USD_PATTERN = re.compile(r"\bUSD\b|\$", re.IGNORECASE)
KZT_PATTERN = re.compile(r"\bKZT\b|тенге|₸|\bтг\b", re.IGNORECASE)


def money(value: float) -> str:
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def main() -> None:
    results = []

    def check(name, passed, detail=None):
        results.append({"check": name, "pass": bool(passed), "detail": detail})

    def near(actual, expected, tolerance=0.01):
        return isinstance(actual, (int, float)) and math.isclose(
            actual, expected, rel_tol=0, abs_tol=tolerance)

    exchange = json.loads((OUT / "evidence" / "exchange-rate.json").read_text(encoding="utf-8"))
    rate = next((exchange[key] for key in ["kzt_per_usd", "rate_kzt_per_usd", "usd_to_kzt", "rate"]
                 if isinstance(exchange.get(key), (int, float))), None)
    if rate is None:
        raise ValueError("exchange-rate.json must include a numeric KZT-per-USD rate")
    check("exchange_rate_positive", rate > 0, {"KZT_per_USD": rate})
    check("exchange_rate_documented", any("http" in str(v) for v in exchange.values()), exchange)

    path = OUT / "SaparTravel_LZ8_Financial_Model.xlsx"
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    errors = [(sh.title, c.coordinate, c.value) for sh in values
              for row in sh for c in row if c.data_type == "e"]
    check("no_spreadsheet_errors", not errors, errors)
    formula_cells = [(sh.title, c.coordinate) for sh in formulas for row in sh for c in row
                     if c.data_type == "f"]
    uncached = [(name, pos) for name, pos in formula_cells if values[name][pos].value is None]
    check("formulas_recalculated", not uncached and len(formula_cells) > 300,
          {"formula_count": len(formula_cells), "uncached": uncached})

    for row in values["Checks"].iter_rows(min_row=2):
        name, actual, expected = (c.value for c in row[:3])
        if name is not None:
            check(f"workbook:{name}", near(actual, expected, 1e-7),
                  {"actual": actual, "expected": expected})

    usd_targets = {
        ("Budget12", "N12"): 29203.80,
        ("Budget12", "N13"): 9600,
        ("Budget12", "N17"): 632,
        ("Budget12", "N18"): 29835.80,
        ("Budget12", "N19"): 160,
        ("Budget12", "N20"): 29995.80,
        ("Hosting", "B12"): 113.65,
        ("Hosting", "C12"): 180.76,
        ("Hosting", "D12"): 211,
        ("Hosting", "B15"): 4131.40,
        ("Hosting", "C15"): 9127.36,
        ("Hosting", "D15"): 7735,
        ("Unit", "B2"): 1.20,
        ("Unit", "B3"): .20,
        ("Unit", "B4"): 1,
        ("Unit", "B5"): 713.65,
        ("Unit", "B6"): 20840,
        ("Unit", "B16"): 3.650475,
        ("Unit", "B18"): 19603.80,
        ("Unit", "B19"): 22735.55,
        ("Unit", "B20"): 23527.55,
    }
    for (sheet, cell), usd in usd_targets.items():
        actual = values[sheet][cell].value
        check(f"KZT:{sheet}!{cell}", near(actual, usd * rate),
              {"actual_KZT": actual, "expected_KZT": usd * rate, "source_USD": usd})

    summary = json.loads((OUT / "assumptions_and_results.json").read_text(encoding="utf-8"))
    check("JSON_currency_KZT", summary.get("currency") == "KZT", summary.get("currency"))
    for key, usd in {"base": 29203.8, "risk": 632, "management": 160,
                     "baseline": 29835.8, "funding": 29995.8, "revenue": 9600,
                     "loss": 19603.8, "fixed": 713.65, "initial": 20840,
                     "contribution": 1, "tco_cloud": 4131.4, "tco_own": 9127.36,
                     "tco_rent": 7735, "peak_deficit": 22735.55}.items():
        actual = summary.get("results", {}).get(key)
        check(f"JSON_results:{key}", near(actual, usd * rate),
              {"actual_KZT": actual, "expected_KZT": usd * rate})

    for sheet, cell, expected in [
        ("Budget12", "N2", 8000), ("Unit", "B8", 714),
        ("Unit", "B9", 1293), ("Unit", "B10", 1315),
        ("Unit", "B11", 2000), ("Unit", "B23", 3068),
    ]:
        actual = values[sheet][cell].value
        check(f"unchanged:{sheet}!{cell}", near(actual, expected, 1e-8),
              {"actual": actual, "expected": expected})

    for row, budgets in [(2, (43.2, 200)), (3, (86.4, 40)), (4, (129.6, 12))]:
        actual = (values["SLO"].cell(row, 5).value, values["SLO"].cell(row, 8).value)
        check(f"C{row-1}_error_budget_unchanged", all(near(a, b, 1e-8) for a, b in zip(actual, budgets)),
              {"actual": actual, "expected": budgets})

    b = values["Budget12"]
    monthly_costs = []
    monthly_revenues = []
    expected_risk_part = float((Decimal("632") * Decimal(str(rate)) / 3).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP))
    previous_cumulative = 0
    for col, users in enumerate(MAU, 2):
        month = col - 1
        usd_lines = [6000 if month <= 3 else 0, 1500 if month == 3 else 0,
                     800 if month == 3 else 0, 500 if month == 4 else 0,
                     40 if month == 1 else 0, 113.65,
                     500 if month >= 4 else 0, 100 if month >= 4 else 0, users * .2]
        expected_cost = sum(usd_lines) * rate
        expected_revenue = users * 1.2 * rate
        actual_lines = [b.cell(row, col).value for row in range(3, 12)]
        cost, revenue = b.cell(12, col).value, b.cell(13, col).value
        monthly_costs.append(cost)
        monthly_revenues.append(revenue)
        check(f"budget_month_{month}_items", all(near(a, u * rate) for a, u in zip(actual_lines, usd_lines)),
              {"actual": actual_lines, "expected": [u * rate for u in usd_lines]})
        check(f"budget_month_{month}_sum", near(sum(actual_lines), cost),
              {"items_sum": sum(actual_lines), "total": cost})
        check(f"budget_month_{month}_converted", near(cost, expected_cost) and near(revenue, expected_revenue),
              {"cost_KZT": cost, "revenue_KZT": revenue})
        check(f"budget_month_{month}_flow", near(b.cell(14, col).value, revenue - cost)
              and near(b.cell(15, col).value, previous_cumulative + revenue - cost),
              {"flow": b.cell(14, col).value, "cumulative": b.cell(15, col).value})
        previous_cumulative = b.cell(15, col).value
        expected_risk = expected_risk_part if month <= 2 else 632 * rate - expected_risk_part * 2 if month == 3 else 0
        check(f"budget_month_{month}_reserves", near(b.cell(17, col).value, expected_risk)
              and near(b.cell(18, col).value, cost + b.cell(17, col).value)
              and near(b.cell(20, col).value, b.cell(18, col).value + b.cell(19, col).value),
              {"risk_KZT": b.cell(17, col).value, "limit_KZT": b.cell(20, col).value})
    for row in range(2, 21):
        if row == 15:
            expected = b.cell(row, 13).value
        elif row == 16:
            expected = b["N12"].value / b["N2"].value
        else:
            expected = sum(b.cell(row, col).value for col in range(2, 14))
        check(f"budget_annual_row_{row}", near(b.cell(row, 14).value, expected),
              {"annual": b.cell(row, 14).value, "computed": expected})

    with (OUT / "budget_12_months.csv").open(encoding="utf-8-sig", newline="") as f:
        csv_all = list(csv.reader(f, delimiter=";"))
    header, csv_rows = csv_all[0], csv_all[1:]
    check("CSV_is_KZT", not USD_PATTERN.search(" ".join(header)) and bool(KZT_PATTERN.search(" ".join(header))), header)
    check("CSV_has_12_months", len(csv_rows) == 12, len(csv_rows))
    for index, row in enumerate(csv_rows):
        numbers = [float(x) for x in row[2:15]]
        check(f"CSV_month_{index+1}_arithmetic", near(sum(numbers[:9]), numbers[9])
              and near(numbers[10] - numbers[9], numbers[11]),
              {"sum": sum(numbers[:9]), "cost": numbers[9], "flow": numbers[11]})
        check(f"CSV_month_{index+1}_matches_workbook", near(numbers[9], monthly_costs[index], .1)
              and near(numbers[10], monthly_revenues[index], .1),
              {"CSV_KZT": numbers[9], "workbook_KZT": monthly_costs[index]})
        if int(row[1]):
            check(f"CSV_month_{index+1}_CPU", near(float(row[15]), numbers[9] / int(row[1]), .01),
                  {"actual": row[15], "computed": numbers[9] / int(row[1])})
    check("CSV_annual_totals_KZT", near(sum(float(row[11]) for row in csv_rows), 29203.8 * rate, 1)
          and near(sum(float(row[12]) for row in csv_rows), 9600 * rate, 1),
          {"costs_KZT": sum(float(row[11]) for row in csv_rows), "revenue_KZT": sum(float(row[12]) for row in csv_rows)})

    check("FX_sheet_matches_evidence", near(values["FX"]["B4"].value, rate, 1e-10), values["FX"]["B4"].value)
    original_tariffs = {2: 24, 3: 30.45, 4: .3, 5: 5, 6: 6, 7: 117.1, 8: 1.9, 9: 59}
    for row, usd in original_tariffs.items():
        check(f"Tariffs_original_and_conversion_row_{row}", near(values["Tariffs"].cell(row, 2).value, usd, 1e-10)
              and near(values["Tariffs"].cell(row, 5).value, usd if row == 4 else usd * rate),
              {"original": values["Tariffs"].cell(row, 2).value, "converted": values["Tariffs"].cell(row, 5).value})

    leaks = []
    for sheet in ["Inputs", "Hosting", "Budget12", "Unit", "Sensitivity"]:
        for row in formulas[sheet]:
            for cell in row:
                if isinstance(cell.value, str) and cell.data_type != "f" and USD_PATTERN.search(cell.value):
                    leaks.append((sheet, cell.coordinate, cell.value))
    check("financial_sheets_use_tenge", not leaks, leaks)
    chart_leaks = []
    with zipfile.ZipFile(path) as archive:
        for filename in archive.namelist():
            if filename.startswith("xl/charts/chart") and filename.endswith(".xml"):
                # Absolute spreadsheet references contain '$'; they are not
                # currency labels. Inspect visible chart text, not formulas.
                chart_xml = ET.fromstring(archive.read(filename))
                chart_text = " ".join(node.text or "" for node in chart_xml.iter()
                                      if node.tag.rsplit("}",1)[-1] == "t")
                if USD_PATTERN.search(chart_text):
                    chart_leaks.append(filename)
    check("financial_charts_use_tenge", not chart_leaks, chart_leaks)

    doc = Document(OUT / "SaparTravel_LZ8_Report.docx")
    table_texts = ["\n".join(c.text for row in table.rows for c in row.cells) for table in doc.tables]
    content = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(table_texts)
    normalized = content.replace("\u00a0", " ").replace("\u202f", " ")
    check("report_contains_required_sections", all(s in content for s in ["Statement of Work", "Error budget", "Юнит-экономика", "Рубежный контроль", "43,2"]),
          {"tables": len(doc.tables), "characters": len(content)})
    check("report_has_KZT_totals", money(29203.8 * rate) in normalized and money(29995.8 * rate) in normalized
          and bool(KZT_PATTERN.search(content)),
          {"annual_KZT": money(29203.8 * rate), "limit_KZT": money(29995.8 * rate)})
    main_table_leaks = []
    for index, (table, text) in enumerate(zip(doc.tables, table_texts), 1):
        first_row = " ".join(c.text for c in table.rows[0].cells)
        # Source evidence may show the provider's original denomination.
        source_table = any(word in first_row.lower() for word in ["источник", "url", "id источника"])
        if not source_table and USD_PATTERN.search(text):
            main_table_leaks.append({"table": index, "header": first_row, "USD_mentions": USD_PATTERN.findall(text)})
    check("report_financial_tables_use_tenge", not main_table_leaks, main_table_leaks)

    prs = Presentation(OUT / "SaparTravel_LZ8_defense.pptx")
    check("defense_has_eight_slides_and_notes", len(prs.slides) == 8 and all(s.has_notes_slide for s in prs.slides), len(prs.slides))
    slide_leaks = []
    slide_content = []
    for index, slide in enumerate(prs.slides, 1):
        visible = "\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        slide_content.append(visible)
        if USD_PATTERN.search(visible):
            slide_leaks.append({"slide": index, "text": visible})
        if USD_PATTERN.search(notes):
            slide_leaks.append({"slide": index, "notes": notes})
    check("presentation_uses_tenge", not slide_leaks, slide_leaks)
    check("presentation_contains_KZT_totals", money(29203.8 * rate) in "\n".join(slide_content)
          and money(29995.8 * rate) in "\n".join(slide_content),
          {"annual_KZT": money(29203.8 * rate), "limit_KZT": money(29995.8 * rate)})

    for name, count in [("backlog_baseline.csv", 18), ("nfr_register.csv", 16)]:
        with (OUT / name).open(encoding="utf-8-sig", newline="") as f:
            records = list(csv.reader(f, delimiter=";"))[1:]
        check(name, len(records) == count, {"rows": len(records), "expected": count})
    for name in ["SaparTravel_LZ8_Report.pdf", "SaparTravel_LZ8_defense.pdf", "evidence/calculator-inputs.json", "evidence/digitalocean-estimate.png"]:
        check(f"deliverable:{name}", (OUT / name).is_file() and (OUT / name).stat().st_size > 0, name)

    with (OUT / "evidence" / "DigitalOcean Estimate Cart.csv").open(encoding="utf-8-sig", newline="") as f:
        provider_rows = list(csv.DictReader(f))
    original_total = sum(float(r["Monthly Price"].replace("$", "")) * int(r["Quantity"]) for r in provider_rows)
    check("official_calculator_original_USD_evidence", len(provider_rows) == 4 and near(original_total, 65.45, 1e-8),
          {"items": len(provider_rows), "original_monthly_USD": original_total, "converted_monthly_KZT": original_total * rate})
    calculator = json.loads((OUT / "evidence" / "calculator-inputs.json").read_text(encoding="utf-8"))
    check("calculator_inputs_original_currency_preserved", near(calculator.get("calculator_base_monthly_usd"), 65.45, 1e-8)
          and near(calculator.get("provider_monthly_usd"), 72.65, 1e-8)
          and calculator.get("paid_resources_created") is False, calculator)

    output = {"as_of": "2026-10-23", "currency": "KZT", "exchange_rate_KZT_per_USD": rate,
              "status": "PASS" if all(r["pass"] for r in results) else "FAIL",
              "scope": "Independent KZT arithmetic, formula caches and document packaging; no production SLO results claimed",
              "formula_count": len(formula_cells), "results": results}
    (OUT / "validation.json").write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": output["status"], "checks": len(results), "formulas": len(formula_cells),
                      "failed": [r for r in results if not r["pass"]]}, ensure_ascii=False, indent=2))
    if output["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
