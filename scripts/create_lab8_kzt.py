"""Build the current LZ8 submission in tenge from the original cost baseline.

Published provider exports are retained in their original currency as evidence.
Run: .venv/bin/python scripts/create_lab8_kzt.py
Then recalculate XLSX and export PDF with LibreOffice before validation/packaging.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import tempfile
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from pptx import Presentation

import create_lab8_package as base


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/lab-8"
RATE = Decimal("447.94")
SOURCE = "https://nationalbank.kz/rss/get_rates.cfm?fdate=23.10.2026"
NUMBER = r"\d+(?:[ \u00a0]\d{3})*(?:[,.]\d+)?"


def kzt(number, places=2):
    value = Decimal(str(number).replace(" ", "").replace("\u00a0", "").replace(",", ".")) * RATE
    return f"{value:,.{places}f}".replace(",", " ").replace(".", ",")


def numeric_cell(text):
    """Convert an already identified monetary cell, including its component sums."""
    return re.sub(NUMBER, lambda m: kzt(m.group()), str(text))


def money_text(text):
    """Only explicitly marked currency amounts; never scale time, SP or MAU."""
    text = str(text)
    def explicit(match):
        value = Decimal(match.group(1).replace(" ", "").replace(",", "."))
        return kzt(match.group(1),4 if 0 < value < 1 else 2) + " тг"
    text = re.sub(r"\$(" + NUMBER + r")", explicit, text)
    text = re.sub(r"(" + NUMBER + r")\s+USD\b", lambda m: kzt(m.group(1)) + " тг", text)
    return text.replace("USD", "тг")


def prose(text):
    text = str(text).replace("Версия: 1.0", "Версия: 1.1 (тенге)")
    if text.startswith("Проект: веб-платформа поиска"):
        return money_text(text) + (
            " Валюта расчетов — казахстанский тенге (KZT). Курс сметы на 23.10.2026: "
            "1 USD = 447,94 тг, официальный RSS Национального Банка РК. Курс зафиксирован "
            "на весь расчетный горизонт; формула пересчета: сумма в тенге = исходная сумма × 447,94. "
            "Денежные таблицы округлены до 0,01 тг, удельные ставки при необходимости показаны "
            "до четырех знаков; промежуточные вычисления хранят полную точность."
        )
    if text.startswith("Генератор:"):
        return ("Генератор текущей версии: scripts/create_lab8_kzt.py; исходная оценка и шаблоны — "
                "scripts/create_lab8_package.py. Запуск: .venv/bin/python scripts/create_lab8_kzt.py. "
                "Затем пересчитать XLSX в LibreOffice, экспортировать PDF и выполнить "
                "scripts/validate_lab8_kzt.py. На листе FX находится курс 447,94 тг за единицу валюты "
                "провайдера; Inputs содержит фиксированные местные затраты в тенге. Provider Tariffs "
                "сохраняет исходные тарифы и формулы пересчета. Изменение курса влияет на иностранные "
                "тарифы, но не индексирует автоматически местную оплату труда и прочие фиксированные "
                "затраты. CSV — контрольный экспорт, XLSX — редактируемая модель.")
    if text.startswith("В проверенном калькуляторе"):
        return ("В калькуляторе переключатель backup описывает weekly-копии и оставлен выключенным. "
                f"Базовый эквивалент {kzt(65.45)} тг подтверждается выгрузкой "
                "evidence/DigitalOcean Estimate Cart.csv, сохраненной в исходной валюте провайдера; "
                f"daily-backup {kzt(7.20)} тг добавлен отдельно по S04. Файл calculator-inputs.json "
                "сохраняет первичные параметры, digitalocean-estimate.png — снимок, "
                "calculator-cart-url.txt — ссылку на корзину. Суммы основной сметы пересчитаны "
                "по единому курсу. Регистрация, создание ресурсов и оплата не выполнялись.")
    if text.startswith("Риск-резерв $632 распределен"):
        total = Decimal(632) * RATE
        first = (total / 3).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        last = total - 2 * first
        f = lambda x: f"{x:,.2f}".replace(",", " ").replace(".", ",")
        return (f"Риск-резерв {f(total)} тг распределен с точностью до тиына: "
                f"{f(first)} + {f(first)} + {f(last)} тг. Труд разработки в WBS оценен "
                f"отдельно от цены подрядчика: 171 ч не умножаются на ставку резерва {kzt(10)} тг/ч. "
                "Расход резерва требует записи события/часов и не считается автоматически понесенным расходом.")
    if text.startswith("В прогнозе первый устойчивый"):
        text = text.replace("29203,80 − 1600", f"{kzt(29203.8)} − {kzt(1600)}")
    if text.startswith("Для 2000 MAU минимальная"):
        text = text.replace("(713,65/2000+0,20)/(500×0,08)",
                            f"({kzt(713.65,3)}/2000+{kzt(.2,3)})/({kzt(500)}×0,08)")
    if text.startswith("Исходный бизнес-кейс ЛЗ3"):
        text += (" В этой версии исторические суммы кейса приведены в эквиваленте тенге по курсу "
                 "сметы; исходные документы недель 1–7 сохранены в архиве как первичные материалы.")
    if text.startswith("Дата проверки официальных веб-источников"):
        text += (" Валюта основной сметы — тенге. Архивные исходники и оригинальная выгрузка "
                 "провайдера могут содержать исходную валюту; их значения не подменяют текущие "
                 "суммы в тенге. Денежные пояснения к тарифам ниже даны в пересчете по курсу сметы.")
    text = text.replace("используются опубликованные USD-цены", "опубликованные цены пересчитываются в тенге по курсу сметы")
    text = text.replace("используются опубликованные USD-тарифы", "опубликованные тарифы пересчитываются в тенге")
    return money_text(text)


class TengeReport(base.Report):
    def __init__(self):
        super().__init__()
        self.doc.sections[0].header.paragraphs[0].text = "SaparTravel / ЛЗ 8 / версия 1.1 / тенге / 23.10.2026"

    def p(self, text):
        super().p(prose(text))

    def h(self, text, level=1):
        super().h(money_text(text), level)

    def bullets(self, items):
        super().bullets([prose(text) for text in items])

    def table(self, headers, rows):
        rows = [[str(c) for c in row] for row in rows]
        if headers[0] == "Результат расчета":
            for n in (0,1,4,5):
                rows[n][1] = numeric_cell(rows[n][1].replace("USD", "тг"))
            rows[2][1] = f"8000 / {kzt(base.make_model()['base']/8000)} тг за MAU-месяц"
        elif headers[0] == "Статья" and "За 12 месяцев, USD" in headers:
            for row in rows:
                row[2] = numeric_cell(row[2])
                if row[0] == "Риск-резерв":
                    row[1] = f"63,2 ч × {kzt(10)} тг/ч; M1/M2 94 366,03; M3 94 366,02 тг"
                if row[0] == "Переменные расходы":
                    row[1] = f"{kzt(.2,4)} тг × 8000 MAU-месяцев"
        elif headers[0] == "Месяц":
            for row in rows:
                for n in (2,3,4,5): row[n] = numeric_cell(row[n]) if row[n] != "N/A" else "N/A"
            rows[-1][5] = kzt(base.make_model()["base"]/8000)
            for row, month in zip(rows, base.make_model()["months"]):
                row[5] = kzt(month["cpu"]) if month["cpu"] is not None else "N/A"
        elif headers[0] == "Переменная":
            for row in rows:
                if row[0] == "Разовые расходы, I": row[1] = numeric_cell(row[1])
                if row[0] == "Распределение I на 36 месяцев": row[1] = f"{kzt(20840)} / 36"
                if row[0] == "Выручка на MAU, r": row[2] = f"{kzt(1.2,4)} тг/мес."
                row[1] = row[1].replace("I+792", f"I+{kzt(792)}")
        elif headers[0] in ("Статья, USD/мес.", "Показатель, USD"):
            for row in rows:
                for n in (1,2,3):
                    if headers[0] == "Статья, USD/мес.":
                        row[n] = re.sub(NUMBER, lambda m:kzt(m.group()), row[n], count=1)
                    else:
                        row[n] = numeric_cell(row[n])
                if row[0].startswith("Общие затраты продукта ("):
                    row[0] = "Общие затраты продукта"
        elif headers[0] == "Тип данных":
            rows[3][1] = ("Проверены 23.10.2026, S01–S08. Для сметы пересчитаны в тенге по курсу "
                          "447,94 тг за единицу валюты провайдера; налоги, инфляция, скидки и free credits не включены.")
        # Values in identified financial cells above have already been converted.
        # Convert explicit currency markers in all other columns exactly once.
        super().table([money_text(x) for x in headers], [[prose(c) for c in row] for row in rows])


def write_documents(model):
    baseline_class = base.Report
    base.Report = TengeReport
    try:
        base.OUT = OUT
        base.write_report(model)
    finally:
        base.Report = baseline_class


def write_slides(source):
    prs = Presentation(source)
    model = base.make_model()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    text = run.text
                    text = text.replace("$3,65 за пользователь-месяц", kzt(model["base"]/8000) + " тг за пользователь-месяц")
                    run.text = prose(text)
        notes = slide.notes_slide.notes_text_frame
        for p in notes.paragraphs:
            for run in p.runs: run.text = prose(run.text)
    # Short currency declaration fits in the subtitle without adding a slide.
    prs.slides[4].shapes[2].text_frame.paragraphs[0].runs[0].text = (
        "09.2026–08.2027 • план в тенге • курс НБРК: 447,94"
    )
    prs.save(OUT / "SaparTravel_LZ8_defense.pptx")


def write_csv(model):
    with (OUT / "budget_12_months.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f,delimiter=";")
        writer.writerow(["Месяц","MAU","Разработка, тг","Миграция, тг","Обучение, тг","Вывод старого, тг","Hosting setup, тг","Размещение, тг","Поддержка, тг","Продвижение, тг","Переменные, тг","Всего, тг","Выручка, тг","Поток, тг","Накопленный поток, тг","тг/MAU"])
        for row in model["months"]:
            scale = lambda x: str((Decimal(str(x))*RATE).quantize(Decimal("0.000001")))
            writer.writerow([row["calendar"],row["mau"],*[scale(x) for x in row["lines"]],
                             scale(row["cost"]),scale(row["revenue"]),scale(row["profit"]),
                             scale(row["cumulative"]),scale(row["cpu"]) if row["cpu"] else "N/A"])


def results_kzt(model):
    invariant={"user_months","operating_be","full_be","reserved_be"}
    result={k:v if k in invariant else float(Decimal(str(v))*RATE)
            for k,v in model.items() if k!="months"}
    result["months"]=[{k:([float(Decimal(str(x))*RATE) for x in v] if k=="lines" else
                          float(Decimal(str(v))*RATE) if k in {"cost","revenue","profit","cumulative","cpu"} and v is not None else v)
                       for k,v in row.items()} for row in model["months"]]
    return result


def package_submission():
    """Package only after PDF export, workbook recalculation and validation."""
    validation=json.loads((OUT/"validation.json").read_text(encoding="utf-8"))
    if validation.get("status")!="PASS" or validation.get("currency")!="KZT":
        raise RuntimeError("Run validate_lab8_kzt.py successfully before packaging")
    files=[p for p in OUT.rglob("*") if p.is_file() and p.suffix!=".zip" and p.name!="manifest.json"]
    files += [ROOT/"README.md",ROOT/"SaparTravel_Charter_V1.0.docx",
              ROOT/"SaparTravel_Business_Case_LZ3.docx",ROOT/"data/throughput.csv",
              ROOT/"notebooks/lab-6-monte-carlo.ipynb"]
    files += [ROOT/"scripts"/name for name in ["create_lab8_package.py","create_lab8_kzt.py",
              "convert_lab8_workbook_kzt.py","validate_lab8_package.py","validate_lab8_kzt.py"]]
    files += [p for p in (ROOT/"docs").rglob("*") if p.is_file() and OUT not in p.parents and not p.name.startswith("~$")]
    files=sorted(set(files))
    assert all(p.is_file() and p.name!=".env" and "node_modules" not in p.parts for p in files)
    manifest={"project":"SaparTravel","assignment":"ЛЗ8 + РК1, недели 1–8","version":"1.1",
              "as_of":"2026-10-23","currency":"KZT","exchange_rate_KZT_per_USD":float(RATE),
              "entry_point":"docs/lab-8/SaparTravel_LZ8_Report.pdf",
              "financial_model":"docs/lab-8/SaparTravel_LZ8_Financial_Model.xlsx",
              "defense":"docs/lab-8/SaparTravel_LZ8_defense.pptx",
              "validation":"docs/lab-8/validation.json",
              "notes":["Основные расчеты ЛЗ8 выражены в казахстанских тенге.",
                       "Исторические материалы недель 1–7 и первичные тарифы сохранены в исходном виде; денежные итоги ЛЗ8 пересчитаны в тенге.",
                       "Расчеты — прогноз, не фактические расходы или подтверждение соблюдения SLA.",
                       "Платные ресурсы не создавались; .env и рабочие секреты в архив не включены."],
              "files":[{"path":str(p.relative_to(ROOT)),"bytes":p.stat().st_size,
                        "sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]}
    mp=OUT/"manifest.json";mp.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    archive=OUT/"SaparTravel_LZ8_Submission.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
        for p in files+[mp]: z.write(p,"SaparTravel_LZ8/"+str(p.relative_to(ROOT)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        for item in manifest["files"]:
            assert hashlib.sha256(z.read("SaparTravel_LZ8/"+item["path"])).hexdigest()==item["sha256"]
    return archive


def main():
    from convert_lab8_workbook_kzt import convert_workbook
    model=base.make_model()
    base.SOURCES.append(("S11","Национальный Банк РК: курс сметы",SOURCE,
                         "Официальный исторический RSS за 23.10.2026, 447,94 тенге за единицу валюты провайдера; исходный XML сохранен в evidence/nbrk-2026-10-23.xml."))
    OUT.mkdir(parents=True,exist_ok=True)
    evidence=OUT/"evidence";evidence.mkdir(exist_ok=True)
    exchange={"date":"2026-10-23","rate":float(RATE),"base":"USD","quote":"KZT","source":SOURCE,
              "status":"official NBRK RSS; quant=1; response archived in nbrk-2026-10-23.xml",
              "policy":"Fixed planning rate; no bank fees, spread or future FX forecast included"}
    (evidence/"exchange-rate.json").write_text(json.dumps(exchange,ensure_ascii=False,indent=2),encoding="utf-8")
    with tempfile.TemporaryDirectory(prefix="sapartravel-lz8-usd-") as temp:
        base.OUT=Path(temp)
        base.write_workbook(model)
        base.write_slides(model)
        convert_workbook(Path(temp)/"SaparTravel_LZ8_Financial_Model.xlsx",OUT/"SaparTravel_LZ8_Financial_Model.xlsx",float(RATE),source_url=SOURCE)
        write_slides(Path(temp)/"SaparTravel_LZ8_defense.pptx")
    write_documents(model)
    write_csv(model)
    converted_inputs=[]
    for label,value,unit,note in base.INPUTS:
        if "USD" in unit: value=float(Decimal(str(value))*RATE)
        converted_inputs.append([label,value,money_text(unit),prose(note)])
    payload={"date":base.DATE,"version":"1.1","currency":"KZT","status":"forecast, not actuals",
             "exchange_rate":exchange,"inputs":converted_inputs,"sources":base.SOURCES,
             "results":results_kzt(model)}
    (OUT/"assumptions_and_results.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"currency":"KZT","rate":float(RATE),"base":payload["results"]["base"],
                      "funding":payload["results"]["funding"],"operating_be":model["operating_be"],"full_be":model["full_be"]},indent=2))


if __name__=="__main__": main()
