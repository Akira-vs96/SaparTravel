"""Validate recalculated LZ8 formulas, independent CSV totals and deliverables.

Run after recalculating XLSX through LibreOffice or Excel:
    .venv/bin/python scripts/validate_lab8_package.py
"""
import csv
import json
import math
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "lab-8"


def main():
    path = OUT / "SaparTravel_LZ8_Financial_Model.xlsx"
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    results = []

    def check(name, passed, detail):
        results.append({"check": name, "pass": bool(passed), "detail": detail})

    errors = [(sh.title, c.coordinate, c.value) for sh in values for row in sh for c in row if c.data_type == "e"]
    check("no_spreadsheet_errors", not errors, errors)
    formula_cells = [(sh.title, c.coordinate) for sh in formulas for row in sh for c in row if c.data_type == "f"]
    uncached = [(name, pos) for name, pos in formula_cells if values[name][pos].value is None]
    check("formulas_recalculated", not uncached and len(formula_cells) > 300,
          {"formula_count": len(formula_cells), "uncached": uncached})
    for row in values["Checks"].iter_rows(min_row=2):
        name, actual, expected = (c.value for c in row[:3])
        check(name, isinstance(actual, (float, int)) and math.isclose(actual, expected, abs_tol=1e-8),
              {"actual": actual, "expected": expected})
    expected = {
        ("Budget12", "N2"): 8000, ("Budget12", "N12"): 29203.80,
        ("Budget12", "N13"): 9600, ("Budget12", "N18"): 29835.80,
        ("Budget12", "N20"): 29995.80, ("Hosting", "B15"): 4131.40,
        ("Hosting", "C15"): 9127.36, ("Hosting", "D15"): 7735,
        ("Unit", "B8"): 714, ("Unit", "B9"): 1293, ("Unit", "B10"): 1315,
        ("Unit", "B16"): 3.650475, ("Unit", "B19"): 22735.55,
        ("Unit", "B20"): 23527.55, ("Unit", "B23"): 3068,
        ("Budget12", "B17"): 210.67, ("Budget12", "C17"): 210.67,
        ("Budget12", "D17"): 210.66,
    }
    for (sheet, cell), target in expected.items():
        actual = values[sheet][cell].value
        check(f"{sheet}!{cell}", isinstance(actual, (float,int)) and math.isclose(actual,target,abs_tol=1e-8),
              {"actual": actual, "expected": target})
    for row, budgets in [(2,(43.2,200)), (3,(86.4,40)), (4,(129.6,12))]:
        actual = (values["SLO"].cell(row,5).value, values["SLO"].cell(row,8).value)
        check(f"C{row-1}_error_budget", all(math.isclose(a,b,abs_tol=1e-8) for a,b in zip(actual,budgets)),
              {"actual": actual, "expected": budgets})
    with (OUT / "budget_12_months.csv").open(encoding="utf-8-sig",newline="") as f:
        rows = list(csv.reader(f,delimiter=";"))[1:]
    costs = sum(float(row[11]) for row in rows)
    revenues = sum(float(row[12]) for row in rows)
    check("CSV_annual_totals", len(rows)==12 and math.isclose(costs,29203.8,abs_tol=1e-8)
          and math.isclose(revenues,9600,abs_tol=1e-8), {"months":len(rows),"costs":costs,"revenue":revenues})
    for m,row in enumerate(rows,2):
        check(f"CSV_month_{m-1}", math.isclose(sum(float(x) for x in row[2:11]),float(row[11]),abs_tol=1e-8)
              and math.isclose(float(row[11]),values["Budget12"].cell(12,m).value,abs_tol=1e-8),row[0])
    for name,count in [("backlog_baseline.csv",18),("nfr_register.csv",16)]:
        with (OUT/name).open(encoding="utf-8-sig",newline="") as f: records=list(csv.reader(f,delimiter=";"))[1:]
        check(name,len(records)==count,{"rows":len(records),"expected":count})
    doc=Document(OUT/"SaparTravel_LZ8_Report.docx")
    content="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    check("report_contains_all_sections", all(s in content for s in ["Statement of Work", "Error budget", "Юнит-экономика", "Рубежный контроль", "43,2", "29 203,80", "29 995,80"]),
          {"tables":len(doc.tables),"characters":len(content)})
    prs=Presentation(OUT/"SaparTravel_LZ8_defense.pptx")
    check("defense_has_eight_slides_and_notes",len(prs.slides)==8 and all(s.has_notes_slide for s in prs.slides),len(prs.slides))
    needed=["SaparTravel_LZ8_Report.pdf","SaparTravel_LZ8_defense.pdf","evidence/calculator-inputs.json","evidence/digitalocean-estimate.png"]
    for name in needed:
        check(f"deliverable:{name}",(OUT/name).is_file() and (OUT/name).stat().st_size>0,name)
    estimate=OUT/"evidence"/"DigitalOcean Estimate Cart.csv"
    if estimate.exists():
        with estimate.open(encoding="utf-8-sig",newline="") as f:
            provider_rows=list(csv.DictReader(f))
        total=sum(float(r["Monthly Price"].replace("$",""))*int(r["Quantity"]) for r in provider_rows)
        check("official_calculator_download",len(provider_rows)==4 and math.isclose(total,65.45,abs_tol=1e-8),
              {"items":len(provider_rows),"base_monthly_usd":total,"daily_backup_separate_usd":7.20})
    else:
        check("official_calculator_download",False,"Missing provider-generated CSV")
    output={"as_of":"2026-10-23","status":"PASS" if all(r["pass"] for r in results) else "FAIL",
            "scope":"Financial arithmetic and document packaging; no application SLO or production tests claimed",
            "results":results}
    (OUT/"validation.json").write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"status":output["status"],"checks":len(results),"formulas":len(formula_cells),
                      "failed":[r for r in results if not r["pass"]]},ensure_ascii=False,indent=2))
    if output["status"]!="PASS": raise SystemExit(1)


if __name__ == "__main__":
    if json.loads((OUT/"assumptions_and_results.json").read_text(encoding="utf-8")).get("currency") == "KZT":
        from validate_lab8_kzt import main as validate_current_package
        validate_current_package()
    else:
        main()
